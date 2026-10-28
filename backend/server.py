from fastapi import FastAPI, APIRouter, HTTPException, Depends, Response, Request, UploadFile, File
from fastapi.security import HTTPBearer
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import logging
from pathlib import Path
from pydantic import BaseModel, Field, EmailStr
from typing import List, Optional
import uuid
from datetime import datetime, timezone, timedelta
import bcrypt
import jwt
import httpx
import io
import re
import hmac
import hashlib
import asyncio
from zoneinfo import ZoneInfo
from zoneinfo import ZoneInfo
from datetime import date as date_type

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

# MongoDB connection
mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

# JWT Settings
JWT_SECRET = os.environ['JWT_SECRET']
JWT_ALGORITHM = 'HS256'
JWT_EXPIRATION_HOURS = 24 * 7  # 7 days

app = FastAPI(title="ISIS Malignani - Metrology Lab Inventory")
api_router = APIRouter(prefix="/api")
security = HTTPBearer(auto_error=False)

# Configure logging
logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

# ============== MODELS ==============

class UserCreate(BaseModel):
    email: EmailStr
    password: str
    name: str

class UserLogin(BaseModel):
    email: EmailStr
    password: str

class UserResponse(BaseModel):
    user_id: str
    email: str
    name: str
    role: str
    picture: Optional[str] = None
    has_badge: bool = False
    created_at: datetime

class LabEntry(BaseModel):
    reason: Optional[str] = None

class LabEntryCreate(BaseModel):
    cognome: str
    nome: str
    class_name: str   # "Prof." oppure la classe dello studente (es. 4A MEC)
    activity: str

class NewsCreate(BaseModel):
    title: str
    body: str
    important: bool = False

class BadgeData(BaseModel):
    badge: str

class UserUpdate(BaseModel):
    name: Optional[str] = None
    role: Optional[str] = None

class InstrumentCreate(BaseModel):
    barcode: str
    matricola: str  # Unique serial number for each instrument
    name: str
    description: Optional[str] = None
    image_url: Optional[str] = None
    location: Optional[str] = None
    calibration_date: Optional[str] = None
    calibration_expiry: Optional[str] = None
    category: Optional[str] = None

class InstrumentUpdate(BaseModel):
    name: Optional[str] = None
    matricola: Optional[str] = None
    description: Optional[str] = None
    image_url: Optional[str] = None
    location: Optional[str] = None
    calibration_date: Optional[str] = None
    calibration_expiry: Optional[str] = None
    category: Optional[str] = None
    status: Optional[str] = None

class InstrumentResponse(BaseModel):
    instrument_id: str
    barcode: str
    matricola: str
    name: str
    description: Optional[str] = None
    image_url: Optional[str] = None
    status: str
    location: Optional[str] = None
    calibration_date: Optional[str] = None
    calibration_expiry: Optional[str] = None
    category: Optional[str] = None
    current_holder: Optional[str] = None
    current_destination: Optional[str] = None
    created_at: datetime

class MovementCreate(BaseModel):
    instrument_id: str
    movement_type: str  # "prelievo" or "deposito"
    destination: Optional[str] = None
    notes: Optional[str] = None

class MovementResponse(BaseModel):
    movement_id: str
    instrument_id: str
    instrument_name: Optional[str] = None
    instrument_matricola: Optional[str] = None
    user_id: str
    user_name: Optional[str] = None
    movement_type: str
    destination: Optional[str] = None
    notes: Optional[str] = None
    timestamp: datetime

# Destination models
class DestinationCreate(BaseModel):
    name: str
    description: Optional[str] = None

class DestinationResponse(BaseModel):
    destination_id: str
    name: str
    description: Optional[str] = None
    created_at: datetime

# ============== HELPER FUNCTIONS ==============

def hash_password(password: str) -> str:
    return bcrypt.hashpw(password.encode('utf-8'), bcrypt.gensalt()).decode('utf-8')

def verify_password(password: str, hashed: str) -> bool:
    return bcrypt.checkpw(password.encode('utf-8'), hashed.encode('utf-8'))

def normalize_badge(raw: str) -> str:
    """Pulisce il codice letto dal lettore: solo lettere e numeri, maiuscole."""
    return re.sub(r"[^A-Za-z0-9]", "", raw or "").upper()

def hash_badge(raw: str) -> str:
    """Il codice della tessera non viene salvato in chiaro: si salva solo un hash con chiave segreta."""
    return hmac.new(JWT_SECRET.encode("utf-8"), normalize_badge(raw).encode("utf-8"), hashlib.sha256).hexdigest()

def public_user(user: dict) -> dict:
    return {k: v for k, v in user.items() if k not in ("password_hash", "badge_hash")} | {"has_badge": bool(user.get("badge_hash"))}

def create_jwt_token(user_id: str, email: str, role: str) -> str:
    payload = {
        'user_id': user_id,
        'email': email,
        'role': role,
        'exp': datetime.now(timezone.utc) + timedelta(hours=JWT_EXPIRATION_HOURS)
    }
    return jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALGORITHM)

def decode_jwt_token(token: str) -> dict:
    try:
        return jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALGORITHM])
    except jwt.ExpiredSignatureError:
        raise HTTPException(status_code=401, detail="Token expired")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Invalid token")

async def get_current_user(request: Request, credentials = Depends(security)) -> dict:
    # First try cookie
    session_token = request.cookies.get('session_token')
    if session_token:
        session = await db.user_sessions.find_one({"session_token": session_token}, {"_id": 0})
        if session:
            expires_at = session.get("expires_at")
            if isinstance(expires_at, str):
                expires_at = datetime.fromisoformat(expires_at)
            if expires_at.tzinfo is None:
                expires_at = expires_at.replace(tzinfo=timezone.utc)
            if expires_at > datetime.now(timezone.utc):
                user = await db.users.find_one({"user_id": session["user_id"]}, {"_id": 0})
                if user:
                    return user
    
    # Then try Authorization header (JWT)
    if credentials and credentials.credentials:
        token = credentials.credentials
        try:
            payload = decode_jwt_token(token)
            user = await db.users.find_one({"user_id": payload["user_id"]}, {"_id": 0})
            if user:
                return user
        except:
            pass
    
    raise HTTPException(status_code=401, detail="Not authenticated")

async def require_admin(current_user: dict = Depends(get_current_user)) -> dict:
    if current_user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="Admin access required")
    return current_user

async def require_admin_or_technician(current_user: dict = Depends(get_current_user)) -> dict:
    if current_user.get("role") not in ["admin", "technician"]:
        raise HTTPException(status_code=403, detail="Admin or Technician access required")
    return current_user

# ============== LOG (registro attività) ==============

async def log_event(user, category: str, action: str, details: str = ""):
    """Scrive una riga nel registro. Non deve mai far fallire l'operazione principale."""
    try:
        user = user or {}
        await db.logs.insert_one({
            "log_id": f"log_{uuid.uuid4().hex[:12]}",
            "timestamp": datetime.now(timezone.utc).isoformat(),
            "user_id": user.get("user_id"),
            "user_name": user.get("name"),
            "user_email": user.get("email"),
            "category": category,
            "action": action,
            "details": details,
        })
    except Exception:
        logger.exception("Impossibile scrivere il log")

# Aule e laboratori dell'ISIS Malignani (aggiunte una sola volta come destinazioni)
SCHOOL_ROOMS = [
    'A.1.01', 'A.1.02', 'A.1.03', 'A.1.04', 'A.1.05', 'A.1.06',
    'A.1.07', 'A.1.08', 'A.1.09', 'A.1.10', 'A.1.11', 'A.1.12',
    'A.1.13', 'A.1.14', 'A.2.01 Lab', 'A.2.02 Lab', 'A.2.03', 'A.2.04',
    'A.2.05', 'A.2.06', 'A.2.07', 'A.2.08', 'A.2.09', 'A.2.10',
    'A.2.11 Lab', 'A.2.12 TEAL', 'A.S.1', 'A.S.2', 'A.S.3', 'A.S.4',
    'B.0.04 Lab', 'B.0.05 CAD', 'B.1.1 ECDL', 'B.2.01 G', 'B.2.02 / B.2.03 Fis G', 'B.2.05',
    'B.2.06', 'B.2.07', 'B.2.10 Lab Chi', 'B.2.12 Lab Chi', 'B.2.15 G', 'B.2.16 G',
    'B.2.17 Lab Fisica', 'B.2.18 b', 'B.2.18 c', 'C.0.01', 'C.0.02', 'C.0.03',
    'C.0.04', 'C.0.05', 'C.0.06', 'C.0.07 TEAL', 'C.0.08 TEAL', 'C.0.09',
    'C.0.10', 'C.1.01', 'C.1.02', 'C.1.03', 'C.1.04', 'C.1.05',
    'C.1.06', 'C.1.07', 'C.1.08', 'C.1.09', 'C.1.10', 'C.2.01',
    'C.2.02', 'C.2.03', 'C.2.04', 'C.2.05', 'C.2.06', 'C.2.07 TEAL',
    'C.2.08 TEAL', 'C.2.09', 'C.2.10', 'E.0.01 Lab', 'E.0.03 G', 'E.0.04 G',
    'E.0.05 SISTEMI ELT', 'E.0.07a Lab SISTEMI', 'E.0.07B', 'E.0.09 G', 'E.0.10 G', 'E.0.11 Lab TDP Tele Eli',
    'E.0.14 Lab ELT', 'E.0.15 MISURE E', 'E.0.16A', 'E.0.16B Lab TDP ELT AUT', 'E.0.18 ASEM Lab', 'E.0.19 Lab Progettazione',
    'E.0.20 Lab', 'E.0.21', 'E.S.03 PC15', 'E.S.14', 'E.S.15', 'E.S.16',
    'E.S.22', 'E.S.23', 'E.S.24', 'E.S.25', 'E.S.26 Aula Immersiva', 'E.S.27',
    'F.0 CNC', 'F.0 Hangar', 'F.0 LAM 1', 'F.0 LAM 2', 'F.0 MEC.01 Metrologia', 'F.0 MEC.02-QMEC',
    'F.0 MEC.03-QTEL', 'F.0 MEC.06', 'F.0 MEC.08', 'F.0 MEC.09', 'F.0 MEC.10', 'F.0 MEC.11 LAB Aggiustaggio',
    'F.0 MEC.11 LAB Carpenteria', 'F.0 MEC.11 LAB Saldatura', 'F.0 MUT.1', 'F.0 MUT.2', 'F.0 MUT.3', 'F.0.02 Lab Prove Mat',
    'F.0.03 G', 'F.0.04 G', 'F.0.06/a Smart Lab', 'F.0.06/b TEAL Smart', 'F.0.08A', 'F.0.08B TEAL',
    'F.0.09 G', 'F.0.10 G', 'F.0.12 - Petri Lab CAD', 'F.0.14', 'F.0.15 TEAL Lardinelli', 'F.0.16',
    'F.0.17 G', 'F.0.19', 'F.0.20 TEAL', 'F.0.22A', 'F.0.22B Lab. TEAL Topografia', 'F.0.23',
    'F.0.24', 'F.0.26', 'F.0.27', 'F.0.28a Lab Mod', 'F.S.05 Lab CAD', 'F.S.10 Lab. chi',
    'F.S.11 aula/F.S.09 Lab. bio', 'F.S.15', 'F.S.17 Lab', 'F.S.18 Lab', 'F.S.19 Lab Inf / CAD', 'F.S.24 Lab CAD',
    'F.S.25', 'LAB MUT', 'LAM 3', 'LAM 4',
]


# ============== AUTH ENDPOINTS ==============

@api_router.post("/auth/register")
async def register(user_data: UserCreate):
    existing = await db.users.find_one({"email": user_data.email}, {"_id": 0})
    if existing:
        raise HTTPException(status_code=400, detail="Email already registered")
    
    user_id = f"user_{uuid.uuid4().hex[:12]}"
    user_count = await db.users.count_documents({})
    
    user_doc = {
        "user_id": user_id,
        "email": user_data.email,
        "name": user_data.name,
        "password_hash": hash_password(user_data.password),
        "role": "admin" if user_count == 0 else "operator",
        "picture": None,
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    
    await db.users.insert_one(user_doc)
    del user_doc["password_hash"]
    user_doc.pop("_id", None)
    await log_event(user_doc, "accessi", "Registrazione", f"Nuovo account creato con ruolo {user_doc['role']}")
    
    token = create_jwt_token(user_id, user_data.email, user_doc["role"])
    
    return {"user": user_doc, "token": token}

@api_router.post("/auth/login")
async def login(user_data: UserLogin, response: Response):
    user = await db.users.find_one({"email": user_data.email}, {"_id": 0})
    if not user or not verify_password(user_data.password, user.get("password_hash", "")):
        await log_event({"email": user_data.email}, "accessi", "Accesso fallito", f"Tentativo con email {user_data.email}")
        raise HTTPException(status_code=401, detail="Invalid credentials")
    
    token = create_jwt_token(user["user_id"], user["email"], user["role"])
    await log_event(user, "accessi", "Accesso", "Login con email e password")
    
    user_response = public_user(user)
    
    return {"user": user_response, "token": token}

@api_router.post("/auth/badge-login")
async def badge_login(data: BadgeData):
    code = normalize_badge(data.badge)
    user = await db.users.find_one({"badge_hash": hash_badge(code)}, {"_id": 0}) if len(code) >= 4 else None
    if not user:
        await asyncio.sleep(0.7)  # rallenta i tentativi a caso
        await log_event(None, "accessi", "Accesso con tessera fallito", "Tessera non riconosciuta")
        raise HTTPException(status_code=401, detail="Tessera non riconosciuta")
    token = create_jwt_token(user["user_id"], user["email"], user["role"])
    await log_event(user, "accessi", "Accesso con tessera", "Login con tessera")
    return {"user": public_user(user), "token": token}

@api_router.post("/auth/google-session")
async def google_session(request: Request, response: Response):
    body = await request.json()
    session_id = body.get("session_id")
    
    if not session_id:
        raise HTTPException(status_code=400, detail="session_id required")
    
    async with httpx.AsyncClient() as client:
        auth_response = await client.get(
            "https://demobackend.emergentagent.com/auth/v1/env/oauth/session-data",
            headers={"X-Session-ID": session_id}
        )
        
        if auth_response.status_code != 200:
            raise HTTPException(status_code=401, detail="Invalid session")
        
        auth_data = auth_response.json()
    
    email = auth_data.get("email")
    user = await db.users.find_one({"email": email}, {"_id": 0})
    
    if not user:
        user_count = await db.users.count_documents({})
        user_id = f"user_{uuid.uuid4().hex[:12]}"
        user = {
            "user_id": user_id,
            "email": email,
            "name": auth_data.get("name", email.split("@")[0]),
            "role": "admin" if user_count == 0 else "operator",
            "picture": auth_data.get("picture"),
            "created_at": datetime.now(timezone.utc).isoformat()
        }
        await db.users.insert_one(user)
        user.pop("_id", None)
    
    session_token = auth_data.get("session_token", f"session_{uuid.uuid4().hex}")
    expires_at = datetime.now(timezone.utc) + timedelta(days=7)
    
    await db.user_sessions.delete_many({"user_id": user["user_id"]})
    await db.user_sessions.insert_one({
        "session_token": session_token,
        "user_id": user["user_id"],
        "expires_at": expires_at.isoformat(),
        "created_at": datetime.now(timezone.utc).isoformat()
    })
    
    response.set_cookie(
        key="session_token",
        value=session_token,
        httponly=True,
        secure=True,
        samesite="none",
        path="/",
        max_age=7 * 24 * 60 * 60
    )
    
    return {"user": user}

@api_router.get("/auth/me")
async def get_me(current_user: dict = Depends(get_current_user)):
    return public_user(current_user)

@api_router.post("/auth/logout")
async def logout(request: Request, response: Response):
    session_token = request.cookies.get("session_token")
    if session_token:
        await db.user_sessions.delete_many({"session_token": session_token})
    
    response.delete_cookie(key="session_token", path="/")
    return {"message": "Logged out successfully"}

# ============== DESTINATIONS ENDPOINTS (Admin manages pickup locations) ==============

@api_router.get("/destinations", response_model=List[DestinationResponse])
async def get_destinations(current_user: dict = Depends(get_current_user)):
    destinations = await db.destinations.find({}, {"_id": 0}).to_list(1000)
    for dest in destinations:
        if isinstance(dest.get("created_at"), str):
            dest["created_at"] = datetime.fromisoformat(dest["created_at"])
    return destinations

@api_router.post("/destinations")
async def create_destination(data: DestinationCreate, current_user: dict = Depends(get_current_user)):
    data.name = (data.name or "").strip()
    if not data.name:
        raise HTTPException(status_code=400, detail="Nome destinazione obbligatorio")
    existing = await db.destinations.find_one(
        {"name": {"$regex": f"^{re.escape(data.name)}$", "$options": "i"}}, {"_id": 0}
    )
    if existing:
        raise HTTPException(status_code=400, detail="Destination already exists")
    
    destination_id = f"dest_{uuid.uuid4().hex[:12]}"
    destination = {
        "destination_id": destination_id,
        "name": data.name,
        "description": data.description,
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    
    await db.destinations.insert_one(destination)
    destination.pop("_id", None)
    await log_event(current_user, "destinazioni", "Destinazione creata", data.name)
    
    return destination

@api_router.put("/destinations/{destination_id}")
async def update_destination(destination_id: str, data: DestinationCreate, current_user: dict = Depends(require_admin)):
    destination = await db.destinations.find_one({"destination_id": destination_id}, {"_id": 0})
    if not destination:
        raise HTTPException(status_code=404, detail="Destination not found")
    
    await db.destinations.update_one(
        {"destination_id": destination_id},
        {"$set": {"name": data.name, "description": data.description}}
    )
    
    updated = await db.destinations.find_one({"destination_id": destination_id}, {"_id": 0})
    await log_event(current_user, "destinazioni", "Destinazione modificata", f"{destination.get('name')} -> {data.name}")
    return updated

@api_router.delete("/destinations/{destination_id}")
async def delete_destination(destination_id: str, current_user: dict = Depends(require_admin)):
    old = await db.destinations.find_one({"destination_id": destination_id}, {"_id": 0})
    result = await db.destinations.delete_one({"destination_id": destination_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Destination not found")
    await log_event(current_user, "destinazioni", "Destinazione eliminata", (old or {}).get("name", destination_id))
    return {"message": "Destination deleted successfully"}

# ============== INSTRUMENTS ENDPOINTS ==============

@api_router.get("/instruments", response_model=List[InstrumentResponse])
async def get_instruments(
    search: Optional[str] = None,
    status: Optional[str] = None,
    category: Optional[str] = None,
    current_user: dict = Depends(get_current_user)
):
    query = {}
    if search:
        query["$or"] = [
            {"barcode": {"$regex": search, "$options": "i"}},
            {"matricola": {"$regex": search, "$options": "i"}},
            {"name": {"$regex": search, "$options": "i"}},
            {"description": {"$regex": search, "$options": "i"}}
        ]
    if status:
        query["status"] = status
    if category:
        query["category"] = category
    
    instruments = await db.instruments.find(query, {"_id": 0}).to_list(1000)
    
    for inst in instruments:
        if isinstance(inst.get("created_at"), str):
            inst["created_at"] = datetime.fromisoformat(inst["created_at"])
        # Ensure matricola field exists for older records
        if "matricola" not in inst:
            inst["matricola"] = inst.get("barcode", "N/A")
    
    return instruments

@api_router.get("/instruments/barcode/{barcode}")
async def get_instruments_by_barcode(barcode: str, current_user: dict = Depends(get_current_user)):
    """Get ALL instruments with this barcode (multiple units with same barcode, different matricola)"""
    instruments = await db.instruments.find({"barcode": barcode}, {"_id": 0}).to_list(1000)
    
    if not instruments:
        raise HTTPException(status_code=404, detail="No instruments found with this barcode")
    
    for inst in instruments:
        if isinstance(inst.get("created_at"), str):
            inst["created_at"] = datetime.fromisoformat(inst["created_at"])
        if "matricola" not in inst:
            inst["matricola"] = inst.get("barcode", "N/A")
    
    return instruments

# ============== EXCEL IMPORT ==============

IMPORT_COLUMNS = [
    ("barcode", "Codice a barre"),
    ("matricola", "Matricola"),
    ("name", "Nome"),
    ("description", "Descrizione"),
    ("category", "Categoria"),
    ("location", "Ubicazione"),
    ("calibration_date", "Data taratura"),
    ("calibration_expiry", "Scadenza taratura"),
    ("image_url", "URL immagine"),
]

# Nomi di colonna accettati (in minuscolo) per ogni campo
IMPORT_ALIASES = {
    "barcode": ["codice a barre", "codice", "barcode", "codice barre"],
    "matricola": ["matricola", "seriale", "numero di serie", "serial"],
    "name": ["nome", "strumento", "name", "nome strumento"],
    "description": ["descrizione", "description"],
    "category": ["categoria", "category"],
    "location": ["ubicazione", "posizione", "location"],
    "calibration_date": ["data taratura", "data calibrazione", "calibration_date", "taratura"],
    "calibration_expiry": ["scadenza taratura", "scadenza calibrazione", "scadenza", "calibration_expiry"],
    "image_url": ["url immagine", "immagine", "image_url", "foto"],
}

def _cell_to_text(value) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date_type):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text or None

def _normalize_date(text: Optional[str]) -> Optional[str]:
    """Converte date tipo 31/12/2026 in 2026-12-31; lascia invariato il resto."""
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return text

@api_router.get("/instruments/import/template")
async def import_template(current_user: dict = Depends(require_admin_or_technician)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Strumenti"
    ws.append([label for _, label in IMPORT_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E3A8A")
    ws.append(["121234567890", "MCEX-001", "Micrometro esterno 0-25 mm", "Esempio di descrizione",
               "Misurazione", "Armadio A", "2026-01-15", "2027-01-15", ""])
    widths = [20, 16, 32, 36, 18, 18, 18, 18, 30]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    # Codice a barre, matricola e date come testo, per non perdere zeri e formato
    for col in ("A", "B", "G", "H"):
        for row in range(2, 502):
            ws[f"{col}{row}"].number_format = "@"
    buffer = io.BytesIO()
    wb.save(buffer)
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="modello_strumenti.xlsx"'},
    )

@api_router.post("/instruments/import/preview")
async def import_preview(file: UploadFile = File(...), current_user: dict = Depends(require_admin_or_technician)):
    from openpyxl import load_workbook
    content = await file.read()
    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="File troppo grande (massimo 5 MB)")
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="Il file non è un Excel valido (.xlsx)")

    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    header = next(rows_iter, None)
    if not header:
        raise HTTPException(status_code=400, detail="Il file è vuoto")

    # Associa le colonne ai campi
    col_index = {}
    for idx, title in enumerate(header):
        t = (str(title).strip().lower() if title is not None else "")
        for field, aliases in IMPORT_ALIASES.items():
            if t in aliases and field not in col_index:
                col_index[field] = idx
    missing = [label for field, label in IMPORT_COLUMNS if field in ("barcode", "matricola", "name") and field not in col_index]
    if missing:
        raise HTTPException(status_code=400, detail="Colonne obbligatorie mancanti: " + ", ".join(missing))

    parsed = []
    for row_number, row in enumerate(rows_iter, start=2):
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue
        if len(parsed) >= 2000:
            raise HTTPException(status_code=400, detail="Troppe righe (massimo 2000)")
        data = {}
        for field, _ in IMPORT_COLUMNS:
            idx = col_index.get(field)
            data[field] = _cell_to_text(row[idx]) if idx is not None and idx < len(row) else None
        data["calibration_date"] = _normalize_date(data["calibration_date"])
        data["calibration_expiry"] = _normalize_date(data["calibration_expiry"])
        parsed.append({"row": row_number, "data": data})
    wb.close()

    existing = set()
    matricole = [r["data"]["matricola"] for r in parsed if r["data"]["matricola"]]
    if matricole:
        async for doc in db.instruments.find({"matricola": {"$in": matricole}}, {"_id": 0, "matricola": 1}):
            existing.add(doc["matricola"])

    seen = set()
    result = []
    for r in parsed:
        d = r["data"]
        status, message = "new", ""
        if not d["barcode"] or not d["matricola"] or not d["name"]:
            status, message = "error", "Mancano codice a barre, matricola o nome"
        elif d["matricola"] in existing:
            status, message = "duplicate", "Matricola già presente nel sistema"
        elif d["matricola"] in seen:
            status, message = "duplicate", "Matricola ripetuta nel file"
        else:
            seen.add(d["matricola"])
        result.append({"row": r["row"], "data": d, "status": status, "message": message})

    return {
        "rows": result,
        "counts": {
            "new": sum(1 for r in result if r["status"] == "new"),
            "duplicate": sum(1 for r in result if r["status"] == "duplicate"),
            "error": sum(1 for r in result if r["status"] == "error"),
        },
    }

class ImportConfirm(BaseModel):
    rows: List[InstrumentCreate]

@api_router.post("/instruments/import/confirm")
async def import_confirm(payload: ImportConfirm, current_user: dict = Depends(require_admin_or_technician)):
    imported, skipped = 0, 0
    seen = set()
    for item in payload.rows:
        if not item.barcode or not item.matricola or not item.name or item.matricola in seen:
            skipped += 1
            continue
        if await db.instruments.find_one({"matricola": item.matricola}, {"_id": 1}):
            skipped += 1
            continue
        seen.add(item.matricola)
        await db.instruments.insert_one({
            "instrument_id": f"inst_{uuid.uuid4().hex[:12]}",
            "barcode": item.barcode,
            "matricola": item.matricola,
            "name": item.name,
            "description": item.description,
            "image_url": item.image_url,
            "status": "available",
            "location": item.location,
            "calibration_date": item.calibration_date,
            "calibration_expiry": item.calibration_expiry,
            "category": item.category,
            "current_holder": None,
            "current_destination": None,
            "created_at": datetime.now(timezone.utc).isoformat(),
        })
        imported += 1
    await log_event(current_user, "strumenti", "Importazione Excel", f"{imported} strumenti importati, {skipped} saltati")
    return {"imported": imported, "skipped": skipped}

@api_router.get("/instruments/{instrument_id}")
async def get_instrument(instrument_id: str, current_user: dict = Depends(get_current_user)):
    instrument = await db.instruments.find_one({"instrument_id": instrument_id}, {"_id": 0})
    if not instrument:
        raise HTTPException(status_code=404, detail="Instrument not found")
    
    if isinstance(instrument.get("created_at"), str):
        instrument["created_at"] = datetime.fromisoformat(instrument["created_at"])
    if "matricola" not in instrument:
        instrument["matricola"] = instrument.get("barcode", "N/A")
    
    return instrument

@api_router.post("/instruments")
async def create_instrument(data: InstrumentCreate, current_user: dict = Depends(require_admin_or_technician)):
    # Check if matricola is unique (matricola must be unique, barcode can be shared)
    existing = await db.instruments.find_one({"matricola": data.matricola}, {"_id": 0})
    if existing:
        raise HTTPException(status_code=400, detail="Matricola already exists")
    
    instrument_id = f"inst_{uuid.uuid4().hex[:12]}"
    instrument = {
        "instrument_id": instrument_id,
        "barcode": data.barcode,
        "matricola": data.matricola,
        "name": data.name,
        "description": data.description,
        "image_url": data.image_url,
        "status": "available",
        "location": data.location,
        "calibration_date": data.calibration_date,
        "calibration_expiry": data.calibration_expiry,
        "category": data.category,
        "current_holder": None,
        "current_destination": None,
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    
    await db.instruments.insert_one(instrument)
    instrument.pop("_id", None)
    await log_event(current_user, "strumenti", "Strumento creato", f"{data.name} (matricola {data.matricola})")
    
    return instrument

@api_router.put("/instruments/{instrument_id}")
async def update_instrument(instrument_id: str, data: InstrumentUpdate, current_user: dict = Depends(require_admin_or_technician)):
    instrument = await db.instruments.find_one({"instrument_id": instrument_id}, {"_id": 0})
    if not instrument:
        raise HTTPException(status_code=404, detail="Instrument not found")
    
    update_data = {k: v for k, v in data.model_dump().items() if v is not None}
    
    # Check matricola uniqueness if being updated
    if "matricola" in update_data:
        existing = await db.instruments.find_one({
            "matricola": update_data["matricola"],
            "instrument_id": {"$ne": instrument_id}
        }, {"_id": 0})
        if existing:
            raise HTTPException(status_code=400, detail="Matricola already exists")
    
    if update_data:
        await db.instruments.update_one({"instrument_id": instrument_id}, {"$set": update_data})
    
    updated = await db.instruments.find_one({"instrument_id": instrument_id}, {"_id": 0})
    changed = ", ".join(sorted(update_data.keys())) or "nessun campo"
    await log_event(current_user, "strumenti", "Strumento modificato", f"{instrument.get('name')} (matricola {instrument.get('matricola')}) - campi: {changed}")
    return updated

@api_router.delete("/instruments/{instrument_id}")
async def delete_instrument(instrument_id: str, current_user: dict = Depends(require_admin_or_technician)):
    old = await db.instruments.find_one({"instrument_id": instrument_id}, {"_id": 0})
    result = await db.instruments.delete_one({"instrument_id": instrument_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Instrument not found")
    await log_event(current_user, "strumenti", "Strumento eliminato", f"{(old or {}).get('name')} (matricola {(old or {}).get('matricola')})")
    
    await db.movements.delete_many({"instrument_id": instrument_id})
    
    return {"message": "Instrument deleted successfully"}

# ============== MOVEMENTS ENDPOINTS ==============

@api_router.post("/movements")
async def create_movement(data: MovementCreate, current_user: dict = Depends(get_current_user)):
    instrument = await db.instruments.find_one({"instrument_id": data.instrument_id}, {"_id": 0})
    if not instrument:
        raise HTTPException(status_code=404, detail="Instrument not found")
    
    if data.movement_type == "prelievo":
        if instrument.get("status") == "in_use":
            raise HTTPException(status_code=400, detail="Instrument is already in use")
        if not data.destination:
            raise HTTPException(status_code=400, detail="Destination is required for prelievo")
        update_data = {
            "status": "in_use",
            "current_holder": current_user["user_id"],
            "current_destination": data.destination
        }
    elif data.movement_type == "deposito":
        if instrument.get("status") == "available":
            raise HTTPException(status_code=400, detail="Instrument is already available")
        update_data = {
            "status": "available",
            "current_holder": None,
            "current_destination": None
        }
    else:
        raise HTTPException(status_code=400, detail="Invalid movement type")
    
    await db.instruments.update_one({"instrument_id": data.instrument_id}, {"$set": update_data})
    
    movement_id = f"mov_{uuid.uuid4().hex[:12]}"
    movement = {
        "movement_id": movement_id,
        "instrument_id": data.instrument_id,
        "instrument_name": instrument.get("name"),
        "instrument_matricola": instrument.get("matricola", instrument.get("barcode")),
        "user_id": current_user["user_id"],
        "user_name": current_user.get("name"),
        "movement_type": data.movement_type,
        "destination": data.destination,
        "notes": data.notes,
        "timestamp": datetime.now(timezone.utc).isoformat()
    }
    
    await db.movements.insert_one(movement)
    movement.pop("_id", None)
    label = "Prelievo" if data.movement_type == "prelievo" else "Deposito"
    where = f" -> {data.destination}" if data.destination else ""
    await log_event(current_user, "movimenti", label, f"{instrument.get('name')} (matricola {instrument.get('matricola', instrument.get('barcode'))}){where}")
    
    return movement

@api_router.get("/movements", response_model=List[MovementResponse])
async def get_movements(
    instrument_id: Optional[str] = None,
    user_id: Optional[str] = None,
    limit: int = 100,
    current_user: dict = Depends(get_current_user)
):
    query = {}
    if instrument_id:
        query["instrument_id"] = instrument_id
    if user_id:
        query["user_id"] = user_id
    
    movements = await db.movements.find(query, {"_id": 0}).sort("timestamp", -1).to_list(limit)
    
    for mov in movements:
        if isinstance(mov.get("timestamp"), str):
            mov["timestamp"] = datetime.fromisoformat(mov["timestamp"])
        if "instrument_matricola" not in mov:
            mov["instrument_matricola"] = None
    
    return movements

@api_router.get("/movements/instrument/{instrument_id}", response_model=List[MovementResponse])
async def get_instrument_movements(instrument_id: str, current_user: dict = Depends(get_current_user)):
    movements = await db.movements.find({"instrument_id": instrument_id}, {"_id": 0}).sort("timestamp", -1).to_list(100)
    
    for mov in movements:
        if isinstance(mov.get("timestamp"), str):
            mov["timestamp"] = datetime.fromisoformat(mov["timestamp"])
        if "instrument_matricola" not in mov:
            mov["instrument_matricola"] = None
    
    return movements

# ============== USERS MANAGEMENT (ADMIN) ==============

@api_router.get("/users", response_model=List[UserResponse])
async def get_users(current_user: dict = Depends(require_admin)):
    users = await db.users.find({}, {"_id": 0, "password_hash": 0}).to_list(1000)
    
    for user in users:
        user["has_badge"] = bool(user.pop("badge_hash", None))
        if isinstance(user.get("created_at"), str):
            user["created_at"] = datetime.fromisoformat(user["created_at"])
    
    return users

@api_router.put("/users/{user_id}")
async def update_user(user_id: str, data: UserUpdate, current_user: dict = Depends(require_admin)):
    user = await db.users.find_one({"user_id": user_id}, {"_id": 0})
    if not user:
        raise HTTPException(status_code=404, detail="User not found")
    
    update_data = {k: v for k, v in data.model_dump().items() if v is not None}
    
    if update_data:
        await db.users.update_one({"user_id": user_id}, {"$set": update_data})
    
    updated = await db.users.find_one({"user_id": user_id}, {"_id": 0, "password_hash": 0, "badge_hash": 0})
    changes = ", ".join(f"{k}={v}" for k, v in update_data.items()) or "nessuna modifica"
    await log_event(current_user, "utenti", "Utente modificato", f"{user.get('name')}: {changes}")
    return updated

@api_router.put("/users/{user_id}/badge")
async def set_user_badge(user_id: str, data: BadgeData, current_user: dict = Depends(require_admin)):
    user = await db.users.find_one({"user_id": user_id}, {"_id": 0})
    if not user:
        raise HTTPException(status_code=404, detail="User not found")
    code = normalize_badge(data.badge)
    if len(code) < 4:
        raise HTTPException(status_code=400, detail="Codice tessera non valido o troppo corto")
    badge_hash = hash_badge(code)
    other = await db.users.find_one({"badge_hash": badge_hash, "user_id": {"$ne": user_id}}, {"_id": 0})
    if other:
        raise HTTPException(status_code=400, detail=f"Tessera già assegnata a {other.get('name', 'un altro utente')}")
    await db.users.update_one({"user_id": user_id}, {"$set": {"badge_hash": badge_hash}})
    await log_event(current_user, "utenti", "Tessera assegnata", f"Utente: {user.get('name')}")
    return {"message": "Tessera assegnata", "has_badge": True}

@api_router.delete("/users/{user_id}/badge")
async def remove_user_badge(user_id: str, current_user: dict = Depends(require_admin)):
    result = await db.users.update_one({"user_id": user_id}, {"$unset": {"badge_hash": ""}})
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="User not found")
    await log_event(current_user, "utenti", "Tessera rimossa", f"user_id: {user_id}")
    return {"message": "Tessera rimossa", "has_badge": False}

@api_router.delete("/users/{user_id}")
async def delete_user(user_id: str, current_user: dict = Depends(require_admin)):
    if user_id == current_user["user_id"]:
        raise HTTPException(status_code=400, detail="Cannot delete yourself")
    
    old = await db.users.find_one({"user_id": user_id}, {"_id": 0})
    result = await db.users.delete_one({"user_id": user_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="User not found")
    
    await db.user_sessions.delete_many({"user_id": user_id})
    await log_event(current_user, "utenti", "Utente eliminato", f"{(old or {}).get('name')} ({(old or {}).get('email')})")
    
    return {"message": "User deleted successfully"}

# ============== REGISTRO ACCESSI AL LABORATORIO ==============

ROME = ZoneInfo("Europe/Rome")

def _today_rome() -> str:
    return datetime.now(ROME).date().isoformat()

async def _open_access(user_id: str):
    """Ultimo ingresso non ancora chiuso da una firma di uscita."""
    return await db.lab_access.find_one(
        {"user_id": user_id, "exit_time": None, "exit_missing": {"$ne": True}},
        {"_id": 0},
        sort=[("entry_time", -1)],
    )

@api_router.get("/lab-access/current")
async def lab_access_current(current_user: dict = Depends(get_current_user)):
    rec = await _open_access(current_user["user_id"])
    if rec and rec.get("date") == _today_rome():
        return {"in_lab": True, "record": rec}
    return {"in_lab": False, "record": None}

@api_router.post("/lab-access/entry")
async def lab_access_entry(data: LabEntry, current_user: dict = Depends(get_current_user)):
    rec = await _open_access(current_user["user_id"])
    if rec:
        if rec.get("date") == _today_rome():
            return {"in_lab": True, "record": rec}
        # ingresso di un giorno passato senza firma di uscita: lo segna come "uscita non firmata"
        await db.lab_access.update_many(
            {"user_id": current_user["user_id"], "exit_time": None},
            {"$set": {"exit_missing": True}},
        )
    now = datetime.now(timezone.utc)
    record = {
        "access_id": f"acc_{uuid.uuid4().hex[:12]}",
        "user_id": current_user["user_id"],
        "user_name": current_user.get("name"),
        "user_email": current_user.get("email"),
        "role": current_user.get("role"),
        "date": now.astimezone(ROME).date().isoformat(),
        "entry_time": now.isoformat(),
        "exit_time": None,
        "exit_missing": False,
        "reason": (data.reason or "").strip()[:200] or None,
    }
    await db.lab_access.insert_one(record)
    record.pop("_id", None)
    await log_event(current_user, "laboratorio", "Ingresso in laboratorio", record["reason"] or "")
    return {"in_lab": True, "record": record}

@api_router.post("/lab-access/exit")
async def lab_access_exit(current_user: dict = Depends(get_current_user)):
    rec = await _open_access(current_user["user_id"])
    if not rec:
        raise HTTPException(status_code=400, detail="Nessun ingresso da chiudere")
    exit_time = datetime.now(timezone.utc).isoformat()
    await db.lab_access.update_one({"access_id": rec["access_id"]}, {"$set": {"exit_time": exit_time}})
    await log_event(current_user, "laboratorio", "Uscita dal laboratorio", "Uscita firmata")
    rec["exit_time"] = exit_time
    return {"in_lab": False, "record": rec}

def _lab_query(date_from: Optional[str], date_to: Optional[str], search: Optional[str]) -> dict:
    query = {}
    dates = {}
    if date_from:
        dates["$gte"] = date_from
    if date_to:
        dates["$lte"] = date_to
    if dates:
        query["date"] = dates
    if search:
        rx = {"$regex": re.escape(search), "$options": "i"}
        query["$or"] = [{"user_name": rx}, {"user_email": rx}, {"reason": rx}]
    return query

@api_router.get("/lab-access")
async def get_lab_access(
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    search: Optional[str] = None,
    limit: int = 100,
    skip: int = 0,
    current_user: dict = Depends(require_admin),
):
    query = _lab_query(date_from, date_to, search)
    total = await db.lab_access.count_documents(query)
    items = await db.lab_access.find(query, {"_id": 0}).sort("entry_time", -1).skip(max(skip, 0)).to_list(max(1, min(limit, 500)))
    return {"items": items, "total": total}

@api_router.get("/lab-access/export")
async def export_lab_access(
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    search: Optional[str] = None,
    current_user: dict = Depends(require_admin),
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    items = await db.lab_access.find(_lab_query(date_from, date_to, search), {"_id": 0}).sort("entry_time", -1).to_list(20000)
    wb = Workbook()
    ws = wb.active
    ws.title = "Accessi laboratorio"
    ws.append(["Data", "Nome", "Email", "Ora ingresso", "Ora uscita", "Durata (min)", "Motivo", "Uscita non firmata"])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E3A8A")
    for it in items:
        entry = datetime.fromisoformat(it["entry_time"]).astimezone(ROME)
        exit_dt = datetime.fromisoformat(it["exit_time"]).astimezone(ROME) if it.get("exit_time") else None
        minutes = round((exit_dt - entry).total_seconds() / 60) if exit_dt else None
        ws.append([
            entry.strftime("%d/%m/%Y"), it.get("user_name") or "-", it.get("user_email") or "-",
            entry.strftime("%H:%M"), exit_dt.strftime("%H:%M") if exit_dt else "",
            minutes if minutes is not None else "", it.get("reason") or "",
            "Sì" if it.get("exit_missing") else "",
        ])
    for col, w in zip("ABCDEFGH", [12, 24, 30, 13, 12, 13, 40, 18]):
        ws.column_dimensions[col].width = w
    buffer = io.BytesIO()
    wb.save(buffer)
    await log_event(current_user, "registro", "Esportazione accessi laboratorio", f"{len(items)} righe")
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="accessi_laboratorio.xlsx"'},
    )

# ============== REGISTRO ACCESSI LABORATORIO ==============

ROME = ZoneInfo("Europe/Rome")

def _rome_day_start_utc(day: str) -> str:
    """'2026-10-28' -> inizio di quel giorno (ora italiana) in UTC, come testo ISO."""
    d = datetime.strptime(day, "%Y-%m-%d").replace(tzinfo=ROME)
    return d.astimezone(timezone.utc).isoformat()

def _lab_query(date_from: Optional[str], date_to: Optional[str], search: Optional[str]) -> dict:
    query = {}
    try:
        if date_from:
            query.setdefault("entry_time", {})["$gte"] = _rome_day_start_utc(date_from)
        if date_to:
            end = datetime.strptime(date_to, "%Y-%m-%d").replace(tzinfo=ROME) + timedelta(days=1)
            query.setdefault("entry_time", {})["$lt"] = end.astimezone(timezone.utc).isoformat()
    except ValueError:
        raise HTTPException(status_code=400, detail="Data non valida (usa AAAA-MM-GG)")
    if search:
        rx = {"$regex": re.escape(search), "$options": "i"}
        query["$or"] = [{"cognome": rx}, {"nome": rx}, {"class_name": rx}, {"activity": rx}]
    return query

async def _open_entry(user_id: str):
    """Ingresso senza uscita. Se è di un giorno passato lo chiude come 'uscita non registrata'."""
    entry = await db.lab_access.find_one(
        {"user_id": user_id, "exit_time": None, "exit_missing": {"$ne": True}}, {"_id": 0}
    )
    if not entry:
        return None
    entered = datetime.fromisoformat(entry["entry_time"]).astimezone(ROME).date()
    if entered < datetime.now(ROME).date():
        await db.lab_access.update_one({"access_id": entry["access_id"]}, {"$set": {"exit_missing": True}})
        return None
    return entry

@api_router.get("/lab/status")
async def lab_status(current_user: dict = Depends(get_current_user)):
    entry = await _open_entry(current_user["user_id"])
    last = await db.lab_access.find({"user_id": current_user["user_id"]}, {"_id": 0}).sort("entry_time", -1).to_list(1)
    last_info = None
    if last:
        last_info = {k: last[0].get(k) for k in ("cognome", "nome", "class_name")}
    return {"inside": bool(entry), "entry": entry, "last": last_info}

@api_router.post("/lab/entry")
async def lab_entry(data: LabEntryCreate, current_user: dict = Depends(get_current_user)):
    cognome, nome = data.cognome.strip(), data.nome.strip()
    class_name, activity = data.class_name.strip(), data.activity.strip()
    if not cognome or not nome or not class_name or not activity:
        raise HTTPException(status_code=400, detail="Compila cognome, nome, classe (o Prof.) e attività svolta")
    existing = await _open_entry(current_user["user_id"])
    if existing:
        return {"inside": True, "entry": existing}
    entry = {
        "access_id": f"acc_{uuid.uuid4().hex[:12]}",
        "user_id": current_user["user_id"],
        "user_email": current_user.get("email"),
        "cognome": cognome[:60],
        "nome": nome[:60],
        "class_name": class_name[:40],
        "activity": activity[:300],
        "entry_time": datetime.now(timezone.utc).isoformat(),
        "exit_time": None,
        "exit_missing": False,
        "signed_by": current_user.get("name"),
    }
    await db.lab_access.insert_one(entry)
    entry.pop("_id", None)
    await log_event(current_user, "laboratorio", "Ingresso in laboratorio", f"{cognome} {nome} ({entry['class_name']}) - {entry['activity']}")
    return {"inside": True, "entry": entry}

@api_router.post("/lab/exit")
async def lab_exit(current_user: dict = Depends(get_current_user)):
    entry = await _open_entry(current_user["user_id"])
    if not entry:
        raise HTTPException(status_code=400, detail="Non risulti dentro il laboratorio")
    exit_time = datetime.now(timezone.utc).isoformat()
    await db.lab_access.update_one({"access_id": entry["access_id"]}, {"$set": {"exit_time": exit_time}})
    entry["exit_time"] = exit_time
    await log_event(current_user, "laboratorio", "Uscita dal laboratorio", f"{entry['cognome']} {entry['nome']}")
    return {"inside": False, "entry": entry}

@api_router.get("/lab/access")
async def lab_access_list(
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    search: Optional[str] = None,
    limit: int = 100,
    skip: int = 0,
    current_user: dict = Depends(require_admin_or_technician),
):
    query = _lab_query(date_from, date_to, search)
    total = await db.lab_access.count_documents(query)
    items = await db.lab_access.find(query, {"_id": 0}).sort("entry_time", -1).skip(max(skip, 0)).to_list(max(1, min(limit, 500)))
    return {"items": items, "total": total}

@api_router.get("/lab/access/export")
async def lab_access_export(
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    search: Optional[str] = None,
    current_user: dict = Depends(require_admin_or_technician),
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    items = await db.lab_access.find(_lab_query(date_from, date_to, search), {"_id": 0}).sort("entry_time", 1).to_list(20000)

    now = datetime.now(ROME)
    start_year = now.year if now.month >= 9 else now.year - 1
    wb = Workbook()
    ws = wb.active
    ws.title = "Registro accessi"
    ws.append(["REGISTRO ACCESSI – LABORATORIO DI METROLOGIA"])
    ws.append([f"Anno scolastico {start_year}/{start_year + 1}"])
    ws.append([])
    headers = ["Data", "Cognome", "Nome", "Classe", "Ora ingresso", "Ora uscita", "Attività svolta", "Firma"]
    ws.append(headers)
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"].font = Font(italic=True)
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in ws[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E3A8A")
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for it in items:
        entered = datetime.fromisoformat(it["entry_time"]).astimezone(ROME)
        exited = datetime.fromisoformat(it["exit_time"]).astimezone(ROME) if it.get("exit_time") else None
        if exited:
            firma = f"Firmato digitalmente: {it.get('signed_by') or ''}"
        elif it.get("exit_missing"):
            firma = f"Uscita non registrata ({it.get('signed_by') or ''})"
        else:
            firma = f"Ingresso firmato: {it.get('signed_by') or ''} (ancora dentro)"
        ws.append([
            entered.strftime("%d/%m/%Y"), it.get("cognome"), it.get("nome"), it.get("class_name"),
            entered.strftime("%H:%M"), exited.strftime("%H:%M") if exited else "", it.get("activity"), firma,
        ])
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    for col, w in zip("ABCDEFGH", [12, 18, 18, 12, 13, 12, 44, 36]):
        ws.column_dimensions[col].width = w
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    buffer = io.BytesIO()
    wb.save(buffer)
    await log_event(current_user, "registro", "Esportazione registro accessi", f"{len(items)} righe")
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="registro_accessi_laboratorio.xlsx"'},
    )

@api_router.get("/lab/access/export-docx")
async def lab_access_export_docx(
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    search: Optional[str] = None,
    current_user: dict = Depends(require_admin_or_technician),
):
    """Registro accessi compilato, nel formato del modulo Word originale."""
    import copy
    from docx import Document
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    template_path = ROOT_DIR / "templates" / "registro_accessi.docx"
    if not template_path.exists():
        raise HTTPException(status_code=500, detail="Modello Word non trovato sul server")

    items = await db.lab_access.find(_lab_query(date_from, date_to, search), {"_id": 0}).sort("entry_time", 1).to_list(5000)

    doc = Document(str(template_path))
    table = doc.tables[0]
    blank_row = copy.deepcopy(table.rows[1]._tr)
    for row in list(table.rows)[1:]:
        table._tbl.remove(row._tr)

    centered = {0, 3, 4, 5}  # Data, Classe, Ora ingresso, Ora uscita

    def add_row(values):
        tr = copy.deepcopy(blank_row)
        table._tbl.append(tr)
        row = table.rows[-1]
        for idx, value in enumerate(values):
            para = row.cells[idx].paragraphs[0]
            if value:
                para.add_run(str(value))
            if idx in centered:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for it in items:
        entered = datetime.fromisoformat(it["entry_time"]).astimezone(ROME)
        exited = datetime.fromisoformat(it["exit_time"]).astimezone(ROME) if it.get("exit_time") else None
        signer = it.get("signed_by") or ""
        if exited:
            firma = f"Firma digitale: {signer}"
        elif it.get("exit_missing"):
            firma = f"{signer} (uscita non firmata)"
        else:
            firma = f"{signer} (ancora dentro)"
        add_row([
            entered.strftime("%d/%m/%Y"), it.get("cognome"), it.get("nome"), it.get("class_name"),
            entered.strftime("%H:%M"), exited.strftime("%H:%M") if exited else "", it.get("activity"), firma,
        ])
    for _ in range(max(0, 12 - len(items))):  # righe vuote come sul foglio cartaceo
        add_row([""] * 8)

    # Anno scolastico nell'intestazione (es. 2026.2027), calcolato dalla data del primo accesso
    ref = datetime.fromisoformat(items[0]["entry_time"]).astimezone(ROME) if items else datetime.now(ROME)
    start_year = ref.year if ref.month >= 9 else ref.year - 1
    for section in doc.sections:
        for node in section.header._element.iter(qn("w:t")):
            if node.text and re.fullmatch(r"\d{4}[./]\d{4}", node.text.strip()):
                node.text = f"{start_year}.{start_year + 1}"

    buffer = io.BytesIO()
    doc.save(buffer)
    await log_event(current_user, "registro", "Esportazione registro accessi (Word)", f"{len(items)} righe")
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": 'attachment; filename="registro_accessi_laboratorio.docx"'},
    )

@api_router.delete("/lab/access/{access_id}")
async def lab_access_delete(access_id: str, current_user: dict = Depends(require_admin)):
    old = await db.lab_access.find_one({"access_id": access_id}, {"_id": 0})
    result = await db.lab_access.delete_one({"access_id": access_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Riga non trovata")
    await log_event(current_user, "laboratorio", "Riga registro accessi eliminata",
                    f"{(old or {}).get('cognome')} {(old or {}).get('nome')} del {(old or {}).get('entry_time', '')[:10]}")
    return {"message": "Riga eliminata"}

# ============== NEWS ==============

@api_router.get("/news")
async def get_news(limit: int = 20, current_user: dict = Depends(get_current_user)):
    items = await db.news.find({}, {"_id": 0}).sort("created_at", -1).to_list(max(1, min(limit, 100)))
    return items

@api_router.post("/news")
async def create_news(data: NewsCreate, current_user: dict = Depends(require_admin_or_technician)):
    title, body = data.title.strip(), data.body.strip()
    if not title or not body:
        raise HTTPException(status_code=400, detail="Titolo e testo sono obbligatori")
    item = {
        "news_id": f"news_{uuid.uuid4().hex[:12]}",
        "title": title[:150],
        "body": body[:3000],
        "important": data.important,
        "author_name": current_user.get("name"),
        "author_id": current_user.get("user_id"),
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.news.insert_one(item)
    item.pop("_id", None)
    await log_event(current_user, "news", "News pubblicata", item["title"])
    return item

@api_router.delete("/news/{news_id}")
async def delete_news(news_id: str, current_user: dict = Depends(require_admin_or_technician)):
    old = await db.news.find_one({"news_id": news_id}, {"_id": 0})
    result = await db.news.delete_one({"news_id": news_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="News non trovata")
    await log_event(current_user, "news", "News eliminata", (old or {}).get("title", news_id))
    return {"message": "News eliminata"}

# ============== LOG (solo admin, sola lettura) ==============

def _logs_query(search: Optional[str], category: Optional[str]) -> dict:
    query = {}
    if category:
        query["category"] = category
    if search:
        rx = {"$regex": re.escape(search), "$options": "i"}
        query["$or"] = [{"details": rx}, {"action": rx}, {"user_name": rx}, {"user_email": rx}]
    return query

@api_router.get("/logs")
async def get_logs(
    search: Optional[str] = None,
    category: Optional[str] = None,
    limit: int = 100,
    skip: int = 0,
    current_user: dict = Depends(require_admin),
):
    query = _logs_query(search, category)
    total = await db.logs.count_documents(query)
    items = await db.logs.find(query, {"_id": 0}).sort("timestamp", -1).skip(max(skip, 0)).to_list(max(1, min(limit, 500)))
    return {"items": items, "total": total}

@api_router.get("/logs/export")
async def export_logs(
    search: Optional[str] = None,
    category: Optional[str] = None,
    current_user: dict = Depends(require_admin),
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from zoneinfo import ZoneInfo
    rome = ZoneInfo("Europe/Rome")
    items = await db.logs.find(_logs_query(search, category), {"_id": 0}).sort("timestamp", -1).to_list(20000)
    wb = Workbook()
    ws = wb.active
    ws.title = "Registro"
    ws.append(["Data", "Ora", "Utente", "Email", "Categoria", "Azione", "Dettagli"])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E3A8A")
    for it in items:
        try:
            dt = datetime.fromisoformat(it["timestamp"]).astimezone(rome)
            d, t = dt.strftime("%d/%m/%Y"), dt.strftime("%H:%M:%S")
        except Exception:
            d, t = it.get("timestamp", ""), ""
        ws.append([d, t, it.get("user_name") or "-", it.get("user_email") or "-", it.get("category"), it.get("action"), it.get("details")])
    for col, w in zip("ABCDEFG", [12, 10, 24, 30, 14, 26, 70]):
        ws.column_dimensions[col].width = w
    buffer = io.BytesIO()
    wb.save(buffer)
    await log_event(current_user, "registro", "Esportazione registro", f"{len(items)} righe")
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="registro_attivita.xlsx"'},
    )

# ============== STATS ENDPOINT ==============

@api_router.get("/stats")
async def get_stats(current_user: dict = Depends(get_current_user)):
    total_instruments = await db.instruments.count_documents({})
    available = await db.instruments.count_documents({"status": "available"})
    in_use = await db.instruments.count_documents({"status": "in_use"})
    total_users = await db.users.count_documents({})
    total_movements = await db.movements.count_documents({})
    
    return {
        "total_instruments": total_instruments,
        "available": available,
        "in_use": in_use,
        "total_users": total_users,
        "total_movements": total_movements
    }

# ============== CATEGORIES ENDPOINT ==============

@api_router.get("/categories")
async def get_categories(current_user: dict = Depends(get_current_user)):
    categories = await db.instruments.distinct("category")
    return [c for c in categories if c]

# Root endpoint
@api_router.get("/")
async def root():
    return {"message": "ISIS Malignani - Metrology Lab Inventory API"}

# Include router
app.include_router(api_router)

# CORS
app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.on_event("startup")
async def seed_school_rooms():
    """Aggiunge una sola volta l'elenco delle aule del Malignani tra le destinazioni."""
    try:
        if await db.meta.find_one({"key": "rooms_seeded_v1"}):
            return
        existing = set()
        async for d in db.destinations.find({}, {"_id": 0, "name": 1}):
            existing.add((d.get("name") or "").strip().lower())
        now = datetime.now(timezone.utc).isoformat()
        docs = [
            {"destination_id": f"dest_{uuid.uuid4().hex[:12]}", "name": r, "description": None, "created_at": now}
            for r in SCHOOL_ROOMS if r.strip().lower() not in existing
        ]
        if docs:
            await db.destinations.insert_many(docs)
        await db.meta.insert_one({"key": "rooms_seeded_v1", "at": now, "added": len(docs)})
        logger.info("Aule aggiunte come destinazioni: %s", len(docs))
    except Exception:
        logger.exception("Impossibile aggiungere le aule")

@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
